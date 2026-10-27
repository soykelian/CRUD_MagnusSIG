from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'docs' / 'planeacion' / 'entrega_0' / 'Entrega_0_Planeacion_MagnusSIS.docx'
LOGO = ROOT / 'assets' / 'branding' / 'logo.png'
BLUE = '0B3F91'
LIGHT = 'EAF2FC'
PALE = 'F4F7FB'
INK = '24364B'
MUTED = '66788F'
WHITE = 'FFFFFF'

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.32)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
for sty,size in [('Title',28),('Heading 1',18),('Heading 2',12),('Heading 3',10)]:
 s=styles[sty]; s.font.name='Aptos Display' if sty in ('Title','Heading 1') else 'Aptos'; s.font.size=Pt(size); s.font.bold=True
 s.font.color.rgb=RGBColor(0,0,0) if sty=='Title' else RGBColor.from_string(BLUE)
 s.paragraph_format.space_before=Pt(7 if sty!='Title' else 0); s.paragraph_format.space_after=Pt(5)
# Remove the built-in Title paragraph rule for a clean cover heading.
pp=styles['Title']._element.get_or_add_pPr()
rule=pp.find(qn('w:pBdr'))
if rule is not None: pp.remove(rule)

def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def margins(cell,top=80,start=100,bottom=80,end=100):
 tc=cell._tc; tcPr=tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
 for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  el=OxmlElement('w:'+tag); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); m.append(el)
 tcPr.append(m)
def set_repeat_table_header(row):
 trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
def set_cell_text(cell,text,bold=False,color=INK,size=8.5):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
 run=p.add_run(text); run.bold=bold; run.font.name='Aptos'; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
def style_table(table,widths=None,header=True,font_size=8.2):
 table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
 if widths:
  for row in table.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 for ri,row in enumerate(table.rows):
  if header and ri==0: set_repeat_table_header(row)
  for cell in row.cells:
   margins(cell)
   if header and ri==0: shade(cell,BLUE)
   else: shade(cell,WHITE if ri%2 else PALE)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    for run in p.runs:
     run.font.name='Aptos'; run.font.size=Pt(font_size)
     if header and ri==0: run.font.bold=True; run.font.color.rgb=RGBColor.from_string(WHITE)
     else: run.font.color.rgb=RGBColor.from_string(INK)

def bullet(text,level=0):
 p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p
def numbered(text):
 p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); p.add_run(text); return p
def label_para(label,text):
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(label+' '); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE); p.add_run(text); return p

# Header and footer
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('MAGNUS SIG  |  DOCUMENTO DE PLANEACIÓN'); r.font.name='Aptos'; r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run('Lenguaje de Programación 1 · 2026-2  |  '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(MUTED)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(10)
p.add_run().add_picture(str(LOGO),width=Inches(1.5))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5); r=p.add_run('Entrega 0 de Proyecto de Aula'); r.font.name='Aptos Display'; r.font.size=Pt(28); r.bold=True; r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
r=p.add_run('Fase de Planeación'); r.font.name='Aptos Display'; r.font.size=Pt(19); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
r=p.add_run('Sistema Web de Gestión de Inventario'); r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(INK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Asignatura: Lenguaje de Programación 1 · 2026-2').font.size=Pt(10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(15)
r=p.add_run('EQUIPO DE TRABAJO'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(BLUE)
names=['Kevin Julian Sanchez Perez','Hilver Galindo','Luis Esteban Medina','Jose Gabriel Cabrera','Brayan Steven Sanchez','Juan Jose Osorio Quitian','Sergio Neira']
t=doc.add_table(rows=4,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for row in t.rows:
 row.cells[0].width=Inches(3.0); row.cells[1].width=Inches(3.0)
for i,name in enumerate(names):
 cell=t.cell(i//2,i%2); set_cell_text(cell,name,size=9); shade(cell,PALE if (i//2)%2==0 else WHITE)
last=t.cell(3,1); set_cell_text(last,'',size=9); shade(last,WHITE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(15)
r=p.add_run('Repositorio del proyecto'); r.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED); r.font.size=Pt(9)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
r=p.add_run('https://github.com/soykelian/CRUD_MagnusSIG'); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(BLUE); r.underline=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Propuesta inicial · Octubre de 2026'); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(MUTED)

doc.add_page_break()
# 1 Project foundation
h=doc.add_heading('1. Fundamentación del proyecto',level=1)
doc.add_heading('Título y descripción',level=2)
doc.add_paragraph('El Sistema Web de Gestión de Inventario Magnus SIG permitirá organizar productos y categorías desde una aplicación web. Busca reducir la dispersión de la información de inventario y facilitar la consulta de existencias, precios y descripciones en un punto común.')
doc.add_heading('Objetivo general',level=2)
doc.add_paragraph('Desarrollar un sistema web de inventario que permita registrar, consultar, buscar, actualizar y eliminar productos, así como organizarlos por categorías y controlar sus cantidades disponibles.')
doc.add_heading('Objetivos específicos',level=2)
for x in [
 'Diseñar e implementar los modelos de productos y categorías con sus relaciones y restricciones de datos.',
 'Construir interfaces web para listar, buscar, registrar, consultar, actualizar y eliminar productos, además de consultar y gestionar categorías.',
 'Aplicar validaciones de entrada para los campos obligatorios, el precio y las cantidades de inventario.',
 'Integrar Django ORM con una base de datos relacional, usando SQLite durante el desarrollo y evaluando MySQL como configuración objetivo.',
 'Verificar las rutas, formularios y operaciones mediante pruebas y documentar la instalación y el uso del sistema.'
]: numbered(x)
doc.add_heading('Alcance',level=2)
for x in [
 'Gestión CRUD de productos: nombre, descripción, categoría, precio y cantidad disponible.',
 'Consulta, búsqueda por nombre o descripción y visualización del detalle de producto.',
 'Listado y administración básica de categorías relacionadas con los productos.',
 'Validación de nombre y categoría obligatorios, precio mayor que cero y cantidad igual o superior a cero.',
 'Persistencia relacional con Django ORM; ambiente local con SQLite y posible configuración de MySQL.'
]: bullet(x)
doc.add_heading('Limitaciones',level=2)
for x in [
 'No incluye facturación, pagos, ventas ni compras.',
 'No contempla reportes avanzados, analítica ni integración con servicios externos.',
 'El control de inventario se limita a la cantidad disponible; no incluye movimientos, lotes, bodegas ni alertas automáticas en esta primera versión.',
 'La autenticación y los permisos por rol se limitan al panel estándar de Django si el tiempo de la asignatura lo permite; no se define un sistema empresarial de usuarios.'
]: bullet(x)

doc.add_page_break()
# 2 Working architecture and roles
doc.add_heading('2. Arquitectura de trabajo y tecnología',level=1)
doc.add_heading('Metodología de trabajo',level=2)
label_para('Metodología propuesta:','Kanban. El equipo organizará el trabajo en un tablero con las columnas Pendiente, En curso, En revisión y Terminado. Cada tarea tendrá responsable y criterio de finalización; se revisará el avance semanalmente y se limitará el trabajo simultáneo para detectar bloqueos temprano.')
label_para('Justificación:','Kanban se adapta a un equipo académico que debe avanzar por entregables y coordinar tareas de interfaz, backend, datos y documentación. Permite visualizar dependencias sin exigir iteraciones Scrum formales cuando las fechas de evaluación dependen del calendario del curso.')
doc.add_heading('Roles y responsabilidades',level=2)
role_rows=[
 ('Frontend','Kevin Julian Sanchez Perez','Maquetar pantallas HTML, CSS y plantillas Django; integrar el prototipo visual y revisar navegación y presentación.'),
 ('Frontend','Hilver Galindo','Apoyar la construcción de interfaces, formularios visuales y adaptación de pantallas a las plantillas del sistema.'),
 ('Backend','Luis Esteban Medina','Desarrollar vistas, rutas y lógica de aplicación en Django; integrar formularios y operaciones de productos.'),
 ('Backend','Jose Gabriel Cabrera','Apoyar modelos, formularios, vistas y operaciones CRUD; coordinar la integración del backend.'),
 ('Base de datos','Brayan Steven Sanchez','Diseñar y mantener el esquema relacional, migraciones y consistencia de los datos.'),
 ('Base de datos','Juan Jose Osorio Quitian','Apoyar configuración de SQLite/MySQL, relaciones ORM y verificación de persistencia.'),
 ('Documentación','Sergio Neira','Mantener documento de requisitos, diagramas, guía de ejecución y registro de decisiones.'),
 ('Documentación / Frontend','Kevin Julian Sanchez Perez','Consolidar la documentación de interfaz y mantenerla alineada con el prototipo y su implementación.')]
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'
for i,v in enumerate(['Área','Responsable','Responsabilidades']): set_cell_text(t.rows[0].cells[i],v,bold=True,color=WHITE,size=8); shade(t.rows[0].cells[i],BLUE)
for row in role_rows:
 cells=t.add_row().cells
 for i,v in enumerate(row): set_cell_text(cells[i],v,size=7.7)
style_table(t,[1.35,2.1,3.4],header=True,font_size=7.7)


doc.add_page_break()
# 3 technology and schedule
doc.add_heading('3. Stack tecnológico y cronograma',level=1)
doc.add_heading('Stack tecnológico',level=2)
tech=[
 ('Lenguajes','Python para backend; HTML5 y CSS3 para la interfaz. JavaScript se reservará para interacciones que lo requieran.'),
 ('Framework','Django con arquitectura MVT, Django Forms y Django ORM para vistas, validación, plantillas y acceso a datos.'),
 ('Base de datos','SQLite para desarrollo local, ya configurado. MySQL es la alternativa objetivo y requiere configuración posterior.'),
 ('Librerías y componentes','Django y sus componentes incluidos. Bootstrap podría evaluarse para acelerar estilos responsivos; el prototipo actual está hecho con HTML y CSS propios y no depende de Bootstrap ni JavaScript.'),
 ('IDE y herramientas','Visual Studio Code como entorno propuesto; Git y GitHub para control de versiones y coordinación.'),
 ('Repositorio','https://github.com/soykelian/CRUD_MagnusSIG')]
t=doc.add_table(rows=1,cols=2); t.style='Table Grid'
for i,v in enumerate(['Componente','Selección y justificación']): set_cell_text(t.rows[0].cells[i],v,bold=True,color=WHITE,size=8); shade(t.rows[0].cells[i],BLUE)
for row in tech:
 cells=t.add_row().cells; set_cell_text(cells[0],row[0],bold=True,size=8); set_cell_text(cells[1],row[1],size=8)
style_table(t,[1.5,5.35],header=True,font_size=8)
doc.add_heading('Cronograma preliminar',level=2)
p=doc.add_paragraph('Fechas tentativas para planear el semestre 2026-2. Deben ajustarse al calendario y a las fechas de evaluación que confirme el profesor.')
p.paragraph_format.space_after=Pt(5)
schedule=[
 ('Planeación y aprobación de requisitos','19–23 oct 2026','Documento de planeación, alcance y backlog inicial'),
 ('Diseño y consolidación del esqueleto','26 oct–6 nov 2026','Diagramas, modelos, rutas y plantillas base'),
 ('CRUD de productos','9–20 nov 2026','Registro, consulta, edición y eliminación'),
 ('Búsqueda, categorías y validaciones','23 nov–4 dic 2026','Búsqueda, gestión de categorías y reglas de datos'),
 ('Integración y pruebas','7–11 dic 2026','Pruebas de rutas, formularios y persistencia'),
 ('Documentación y entrega final','14–18 dic 2026','README, manual breve y versión revisada')]
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'
for i,v in enumerate(['Fase','Fechas tentativas','Resultado esperado']): set_cell_text(t.rows[0].cells[i],v,bold=True,color=WHITE,size=8); shade(t.rows[0].cells[i],BLUE)
for row in schedule:
 cells=t.add_row().cells
 for i,v in enumerate(row): set_cell_text(cells[i],v,size=7.7)
style_table(t,[2.1,1.45,3.3],header=True,font_size=7.7)
# Add a compact Gantt view
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3)
r=p.add_run('Vista de fases'); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
months=['Fase','Oct 19–30','Nov 2–20','Nov 23–Dec 4','Dec 7–18']
rows=[months,
 ['Planeación','●','','',''],
 ['Esqueleto','●','●','',''],
 ['CRUD productos','','●','',''],
 ['Categorías y validaciones','','','●',''],
 ['Integración y cierre','','','','●']]
t=doc.add_table(rows=0,cols=5); t.style='Table Grid'
for ri,row in enumerate(rows):
 cells=t.add_row().cells
 for i,v in enumerate(row):
  set_cell_text(cells[i],v,bold=(ri==0),color=WHITE if ri==0 else (BLUE if v=='●' else INK),size=7.4)
  if ri==0: shade(cells[i],BLUE)
  elif v=='●': shade(cells[i],LIGHT)
  else: shade(cells[i],WHITE if ri%2 else PALE)
style_table(t,[1.75,1.1,1.1,1.4,1.5],header=True,font_size=7.4)

# Document properties
doc.core_properties.title='Entrega 0 de Proyecto de Aula - Fase de Planeación'
doc.core_properties.subject='Sistema Web de Gestión de Inventario Magnus SIG'
doc.core_properties.author='Equipo de proyecto Magnus SIG'
doc.core_properties.keywords='planeación, inventario, Django, Kanban'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print('tables',len(doc.tables),'paragraphs',len(doc.paragraphs))


